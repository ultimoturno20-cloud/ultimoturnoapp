from __future__ import annotations

import csv
import html
import json
import re
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = ROOT / "outputs" / "old_claim_import"

CLAIM_SOURCES = [
    {
        "date": "2026-07-12",
        "reference": "Claim 12 de julio",
        "path": Path(r"C:\Users\skype\Downloads\Claim 12 de Julio.xlsx"),
        "url": "https://docs.google.com/spreadsheets/d/1Y7pFN6lspt2gLPgb9k1IjiPCLzIv0Oc5KpERjXQjR7M/edit?gid=671161712#gid=671161712",
        "extra_sheets": [
            {"name": "GRATIS Y REPES 12 JULIO", "date": "2026-07-12", "reference": "Claim 12 de julio"},
            {"name": "free y repes 2 julio", "date": "2026-07-02", "reference": "Claim 2 de julio"},
        ],
    },
    {
        "date": "2026-07-26",
        "reference": "Claim 26 de julio",
        "path": Path(r"C:\Users\skype\Downloads\Claim 26 de julio.xlsx"),
        "url": "https://docs.google.com/spreadsheets/d/1ls7fZ5u90FWWGSuJeAz7NOmLruT6o2EUXRLDZPzFeio/edit?gid=250762855#gid=250762855",
        "extra_sheets": [
            {"name": "gratis y repetidas", "date": "2026-07-26", "reference": "Claim 26 de julio"},
        ],
    },
]

ADMIN_SOURCE = {
    "path": Path(r"C:\Users\skype\Downloads\Administracion Melody.xlsx"),
    "url": "https://docs.google.com/spreadsheets/d/1nZA4qvZZk6a0dQfkPY8dVIAc30ISIt6pacOkOI-dspI/edit?gid=2096430480#gid=2096430480",
}

ORDER_HEADERS = [
    "Order ID",
    "Referencia",
    "Fecha referencia",
    "Comprador",
    "Total ARS",
    "Total USD",
    "Cartas",
    "Pagado",
    "Embalado",
    "Entregado",
    "Fecha pago",
    "Fecha entrega",
    "Sync ventas",
    "Referencia URL",
    "Notas",
]

SALES_HEADERS = [
    "Venta ID",
    "Order ID",
    "Referencia",
    "Fecha referencia",
    "Fecha venta",
    "Origen",
    "Comprador",
    "Nombre final",
    "Nombre",
    "Expansion",
    "Cantidad",
    "Precio ARS",
    "Precio USD",
    "PriceCharting URL",
    "PriceCharting ID",
    "SKU",
    "Pagado",
    "Embalado",
    "Entregado",
    "Sync Stock",
    "Tags",
    "Notas",
]

FREE_HEADERS = [
    "Free ID",
    "Order ID",
    "Referencia",
    "Fecha referencia",
    "Comprador",
    "Nombre final",
    "Nombre",
    "Expansion",
    "Cantidad",
    "PriceCharting URL",
    "PriceCharting ID",
    "SKU",
    "Entregado",
    "Sync Stock",
    "Tags",
    "Notas",
]

REVIEW_HEADERS = [
    "Tipo",
    "Fuente",
    "Hoja",
    "Fila",
    "Fecha",
    "Comprador",
    "Dato",
    "ARS",
    "USD",
    "Notas",
]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    return html.unescape(str(value)).replace("\n", " ").strip()


def normalize_text(value: Any) -> str:
    text = clean_text(value).lower()
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def slug(value: Any) -> str:
    text = normalize_text(value)
    return re.sub(r"[^a-z0-9]+", "-", text).strip("-")


def buyer_key(value: Any) -> str:
    text = normalize_text(value)
    digits = re.findall(r"\d", clean_text(value))
    last4 = "".join(digits)[-4:] if len(digits) >= 4 else ""
    return f"{text}|{last4}" if last4 else text


def money(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = clean_text(value).lower()
    text = re.sub(r"(ars|usd|\$)", "", text).strip()
    if not text:
        return 0.0
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def parse_date(value: Any, fallback: str) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = clean_text(value)
    if re.match(r"^\d{4}-\d{2}-\d{2}", text):
        return text[:10]
    return fallback


def parse_price(final_price: Any, pc_usd: Any, final_name: str) -> tuple[float, float, str]:
    price = money(final_price)
    usd_ref = money(pc_usd)
    notes = []
    if not price:
        return 0.0, 0.0, "Sin precio final"
    lower = final_name.lower()
    if "usd" in lower or (price < 500 and usd_ref >= 20):
        notes.append("Precio final viejo interpretado como USD")
        return 0.0, price, "; ".join(notes)
    return price, 0.0, ""


def pc_id_from_url(url: str) -> str:
    match = re.search(r"[?&]id=(\d+)", url or "")
    return match.group(1) if match else ""


def header_map(ws) -> dict[str, int]:
    out = {}
    for index, cell in enumerate(ws[1], start=1):
        header = normalize_text(cell.value)
        if header:
            out[header] = index
    return out


def get_cell(ws, row: int, col: int | None) -> Any:
    return ws.cell(row, col).value if col else None


def add_review(reviews: list[list[Any]], kind: str, source: str, sheet: str, row: int, date: str, buyer: str, data: str, ars: float, usd: float, notes: str) -> None:
    reviews.append([kind, source, sheet, row or "", date, buyer, data, ars or "", usd or "", notes])


def read_admin_statuses(reviews: list[list[Any]]) -> dict[str, list[dict[str, Any]]]:
    statuses: dict[str, list[dict[str, Any]]] = defaultdict(list)
    path = ADMIN_SOURCE["path"]
    if not path.exists():
        add_review(reviews, "ERROR", path.name, "", "", "", "", "", 0, 0, "No existe Administracion Melody.xlsx")
        return statuses

    wb = load_workbook(path, data_only=True)
    if "Hoja 19" in wb.sheetnames:
        ws = wb["Hoja 19"]
        blocks = [(1, 2, 3, 4, 5), (8, 9, 10, 11, None)]
        for name_col, total_col, usd_col, date_col, paid_col in blocks:
            for row in range(2, ws.max_row + 1):
                name = clean_text(ws.cell(row, name_col).value)
                if not name:
                    continue
                total_ars = money(ws.cell(row, total_col).value)
                total_usd = money(ws.cell(row, usd_col).value)
                raw_date = ws.cell(row, date_col).value
                date = parse_date(raw_date, "2026-07-26")
                paid_value = ws.cell(row, paid_col).value if paid_col else None
                paid = paid_value is True
                key = f"{date}|{buyer_key(name)}"
                statuses[key].append({
                    "source": "Administracion Melody / Hoja 19",
                    "row": row,
                    "buyer": name,
                    "ars": total_ars,
                    "usd": total_usd,
                    "date": date,
                    "paid": paid,
                    "rawPaid": clean_text(paid_value),
                    "notes": clean_text(raw_date) if not re.match(r"^\d{4}-\d{2}-\d{2}", clean_text(raw_date)) else "",
                })

    if "a cobrar" in wb.sheetnames:
        ws = wb["a cobrar"]
        for row in range(2, ws.max_row + 1):
            name = clean_text(ws.cell(row, 1).value)
            total_ars = money(ws.cell(row, 2).value)
            total_usd = money(ws.cell(row, 3).value)
            if name and (total_ars or total_usd):
                add_review(reviews, "ADMIN_DEUDA_MANUAL", path.name, "a cobrar", row, "", name, "Fila manual sin detalle carta por carta", total_ars, total_usd, "No se importa automaticamente")

    return statuses


def source_header(source: dict[str, Any]) -> str:
    return source["path"].name


def read_claim_sales(source: dict[str, Any], reviews: list[list[Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    path = source["path"]
    wb = load_workbook(path, data_only=True)
    ws = wb["LISTA"]
    headers = header_map(ws)
    cols = {
        "nombre": headers.get("nombre"),
        "expansion": headers.get("expansion"),
        "link": headers.get("link"),
        "usd": headers.get("usd"),
        "precio_final": headers.get("precio final"),
        "nombre_final": headers.get("nombrefinal") or headers.get("nombre final"),
        "comprador": headers.get("comprador"),
        "tags": headers.get("tags"),
    }

    sales: list[dict[str, Any]] = []
    suspicious: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        buyer = clean_text(get_cell(ws, row, cols["comprador"]))
        final_name = clean_text(get_cell(ws, row, cols["nombre_final"]))
        name = clean_text(get_cell(ws, row, cols["nombre"]))
        expansion = clean_text(get_cell(ws, row, cols["expansion"]))
        pc_url = clean_text(get_cell(ws, row, cols["link"]))
        if not buyer or not (final_name or name):
            continue

        ars, usd, price_note = parse_price(get_cell(ws, row, cols["precio_final"]), get_cell(ws, row, cols["usd"]), final_name)
        if not ars and not usd:
            continue

        pc_id = pc_id_from_url(pc_url)
        notes = [f"Import viejo: {path.name} / LISTA fila {row}"]
        if price_note:
            notes.append(price_note)
            suspicious.append({
                "row": row,
                "buyer": buyer,
                "name": final_name or name,
                "ars": ars,
                "usd": usd,
                "note": price_note,
            })

        sales.append({
            "source": path.name,
            "sheet": "LISTA",
            "row": row,
            "reference": source["reference"],
            "date": source["date"],
            "url": source["url"],
            "buyer": buyer,
            "buyerKey": buyer_key(buyer),
            "finalName": final_name or build_final_name(name, expansion, ars, usd),
            "name": name,
            "expansion": expansion,
            "quantity": 1,
            "ars": ars,
            "usd": usd,
            "pcUrl": pc_url,
            "pcId": pc_id,
            "sku": f"PKM-PC-{pc_id}" if pc_id else "",
            "tags": clean_text(get_cell(ws, row, cols["tags"])),
            "notes": " | ".join(notes),
        })

    for item in suspicious:
        add_review(reviews, "PRECIO_USD_DETECTADO", path.name, "LISTA", item["row"], source["date"], item["buyer"], item["name"], item["ars"], item["usd"], item["note"])
    return sales, suspicious


def build_final_name(name: str, expansion: str, ars: float, usd: float) -> str:
    base = " - ".join(part for part in [name, expansion] if part)
    parts = []
    if ars:
        parts.append(f"${int(ars)}")
    if usd:
        parts.append(f"${usd:g}usd")
    return f"{base} - {' + '.join(parts)}" if parts else base


def read_extra_sheets(source: dict[str, Any], reviews: list[list[Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    path = source["path"]
    wb = load_workbook(path, data_only=True)
    extra_sales: list[dict[str, Any]] = []
    frees: list[dict[str, Any]] = []

    for extra in source.get("extra_sheets", []):
        sheet_name = extra["name"]
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        date = extra["date"]
        reference = extra["reference"]

        for row in range(2, ws.max_row + 1):
            free_name = clean_text(ws.cell(row, 1).value)
            free_buyer = clean_text(ws.cell(row, 2).value)
            if free_name and free_buyer:
                frees.append({
                    "source": path.name,
                    "sheet": sheet_name,
                    "row": row,
                    "reference": reference,
                    "date": date,
                    "url": source["url"],
                    "buyer": free_buyer,
                    "buyerKey": buyer_key(free_buyer),
                    "finalName": free_name,
                    "name": free_name,
                    "expansion": "",
                    "quantity": 1,
                    "pcUrl": "",
                    "pcId": "",
                    "sku": "",
                    "tags": "",
                    "notes": f"Free importado desde {path.name} / {sheet_name} fila {row}",
                })

            rep_name = clean_text(ws.cell(row, 4).value)
            rep_price_raw = ws.cell(row, 5).value
            rep_buyer = clean_text(ws.cell(row, 6).value)
            if rep_name and rep_buyer:
                raw_price_text = clean_text(rep_price_raw).lower()
                price_value = money(rep_price_raw)
                is_usd = "usd" in raw_price_text
                ars = 0.0 if is_usd else price_value
                usd = price_value if is_usd else 0.0
                extra_sales.append({
                    "source": path.name,
                    "sheet": sheet_name,
                    "row": row,
                    "reference": reference,
                    "date": date,
                    "url": source["url"],
                    "buyer": rep_buyer,
                    "buyerKey": buyer_key(rep_buyer),
                    "finalName": rep_name,
                    "name": rep_name,
                    "expansion": "",
                    "quantity": 1,
                    "ars": ars,
                    "usd": usd,
                    "pcUrl": "",
                    "pcId": "",
                    "sku": "",
                    "tags": "",
                    "notes": f"Repetida/venta extra importada desde {path.name} / {sheet_name} fila {row}",
                })
            elif rep_name or rep_buyer:
                add_review(reviews, "EXTRA_INCOMPLETA", path.name, sheet_name, row, date, rep_buyer, rep_name, money(rep_price_raw), 0, "Fila derecha sin carta o comprador claro")

    return extra_sales, frees


def make_order_id(date: str, buyer: str, used: set[str]) -> str:
    base = f"{date.replace('-', '')}-{slug(buyer).upper()[:42]}"
    order_id = base
    index = 2
    while order_id in used:
        order_id = f"{base}-{index}"
        index += 1
    used.add(order_id)
    return order_id


def match_admin_status(date: str, buyer: str, statuses: dict[str, list[dict[str, Any]]]) -> dict[str, Any] | None:
    exact = statuses.get(f"{date}|{buyer_key(buyer)}")
    if exact:
        return exact[0]
    digits = re.findall(r"\d", clean_text(buyer))
    last4 = "".join(digits)[-4:] if len(digits) >= 4 else ""
    if not last4:
        return None
    for key, values in statuses.items():
        if key.startswith(f"{date}|") and key.endswith(f"|{last4}"):
            return values[0]
    return None


def build_import() -> dict[str, Any]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    reviews: list[list[Any]] = []
    admin_statuses = read_admin_statuses(reviews)

    sales_items: list[dict[str, Any]] = []
    free_items: list[dict[str, Any]] = []
    for source in CLAIM_SOURCES:
        main_sales, _ = read_claim_sales(source, reviews)
        extra_sales, frees = read_extra_sheets(source, reviews)
        sales_items.extend(main_sales)
        sales_items.extend(extra_sales)
        free_items.extend(frees)

    grouped_sales: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for item in sales_items:
        grouped_sales[(item["date"], item["buyerKey"])].append(item)

    grouped_frees: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for item in free_items:
        grouped_frees[(item["date"], item["buyerKey"])].append(item)

    orders: list[list[Any]] = []
    sales_rows: list[list[Any]] = []
    free_rows: list[list[Any]] = []
    order_lookup: dict[tuple[str, str], str] = {}
    used_order_ids: set[str] = set()
    summary_rows: list[list[Any]] = []

    all_keys = sorted(set(grouped_sales) | set(grouped_frees), key=lambda x: (x[0], x[1]))
    for date, key in all_keys:
        sale_group = grouped_sales.get((date, key), [])
        free_group = grouped_frees.get((date, key), [])
        first = (sale_group or free_group)[0]
        buyer = first["buyer"]
        reference = first["reference"]
        source_url = first["url"]
        order_id = make_order_id(date, buyer, used_order_ids)
        order_lookup[(date, key)] = order_id
        total_ars = round(sum(float(item.get("ars") or 0) for item in sale_group), 2)
        total_usd = round(sum(float(item.get("usd") or 0) for item in sale_group), 2)
        cards = sum(int(item.get("quantity") or 1) for item in sale_group)
        status = match_admin_status(date, buyer, admin_statuses)
        paid = bool(status and status.get("paid"))
        notes = [f"Migrado desde claims viejos; frees: {len(free_group)}"]
        if status:
            notes.append(f"Estado Admin: {status['source']} fila {status['row']} paid={status.get('rawPaid')}")
            admin_ars = round(float(status.get("ars") or 0), 2)
            admin_usd = round(float(status.get("usd") or 0), 2)
            if admin_ars and abs(admin_ars - total_ars) > 5:
                add_review(reviews, "TOTAL_DISTINTO_ADMIN", status["source"], "", status["row"], date, buyer, "Total claim vs admin", total_ars, total_usd, f"Admin ARS={admin_ars} USD={admin_usd}")
        else:
            add_review(reviews, "SIN_MATCH_ADMIN", "Claims import", "", "", date, buyer, reference, total_ars, total_usd, "No encontre comprador+fecha en Administracion Melody")

        orders.append([
            order_id,
            reference,
            date,
            buyer,
            total_ars or "",
            total_usd or "",
            cards,
            paid,
            False,
            False,
            "",
            "",
            "",
            source_url,
            " | ".join(notes),
        ])

        summary_rows.append([
            date,
            reference,
            buyer,
            order_id,
            cards,
            len(free_group),
            total_ars or "",
            total_usd or "",
            "Si" if paid else "No",
            "Si" if status else "No",
        ])

        for idx, item in enumerate(sale_group, start=1):
            sales_rows.append([
                f"{order_id}-{idx}",
                order_id,
                item["reference"],
                item["date"],
                "",
                "Claim",
                item["buyer"],
                item["finalName"],
                item["name"],
                item["expansion"],
                item["quantity"],
                item["ars"] or "",
                item["usd"] or "",
                item["pcUrl"],
                item["pcId"],
                item["sku"],
                paid,
                False,
                False,
                "",
                item.get("tags", ""),
                item["notes"],
            ])

        for idx, item in enumerate(free_group, start=1):
            free_rows.append([
                f"{order_id}-FREE-{idx}",
                order_id,
                item["reference"],
                item["date"],
                item["buyer"],
                item["finalName"],
                item["name"],
                item["expansion"],
                item["quantity"],
                item["pcUrl"],
                item["pcId"],
                item["sku"],
                False,
                "",
                item.get("tags", ""),
                item["notes"],
            ])

    matched_admin_keys = {f"{date}|{key}" for date, key in order_lookup}
    for key, values in admin_statuses.items():
        if key in matched_admin_keys:
            continue
        for status in values:
            add_review(reviews, "ADMIN_SIN_DETALLE", status["source"], "Hoja 19", status["row"], status["date"], status["buyer"], "Resumen sin cartas en claims cargados", status["ars"], status["usd"], status.get("notes", ""))

    data = {
        "metadata": {
            "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "sources": [str(source["path"]) for source in CLAIM_SOURCES] + [str(ADMIN_SOURCE["path"])],
        },
        "headers": {
            "orders": ORDER_HEADERS,
            "sales": SALES_HEADERS,
            "frees": FREE_HEADERS,
            "review": REVIEW_HEADERS,
            "summary": ["Fecha", "Referencia", "Comprador", "Order ID", "Cartas", "Frees", "Total ARS", "Total USD", "Pagado sugerido", "Match admin"],
        },
        "orders": orders,
        "sales": sales_rows,
        "frees": free_rows,
        "review": reviews,
        "summary": summary_rows,
    }
    return data


def write_csv(path: Path, headers: list[str], rows: list[list[Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(rows)


def main() -> None:
    data = build_import()
    json_path = OUTPUT_DIR / "old_claim_import.json"
    json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    write_csv(OUTPUT_DIR / "ordenes_import.csv", data["headers"]["orders"], data["orders"])
    write_csv(OUTPUT_DIR / "ventas_detalle_import.csv", data["headers"]["sales"], data["sales"])
    write_csv(OUTPUT_DIR / "frees_import.csv", data["headers"]["frees"], data["frees"])
    write_csv(OUTPUT_DIR / "revision_import.csv", data["headers"]["review"], data["review"])
    print(json.dumps({
        "orders": len(data["orders"]),
        "sales": len(data["sales"]),
        "frees": len(data["frees"]),
        "review": len(data["review"]),
        "output": str(json_path),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
