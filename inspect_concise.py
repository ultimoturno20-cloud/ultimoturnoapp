from pathlib import Path
import openpyxl


BASE = Path(r"C:\Users\skype\Downloads\New folder")
FILES = [
    "Presupuestador.xlsx",
    "extraer imagen.xlsx",
    "Generar imagen adelanto.xlsx",
    "Claim 12 de Julio.xlsx",
]


def clean(value):
    if value is None:
        return ""
    if hasattr(value, "text"):
        value = value.text
    text = str(value).replace("\n", "\\n")
    return text if len(text) <= 100 else text[:97] + "..."


def formula_samples(ws, limit=10):
    out = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                out.append(f"{cell.coordinate}: {clean(cell.value)}")
                if len(out) >= limit:
                    return out
    return out


def hyperlink_count(ws):
    count = 0
    samples = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.hyperlink:
                count += 1
                if len(samples) < 5:
                    samples.append(f"{cell.coordinate}: {cell.hyperlink.target}")
    return count, samples


def print_sheet(ws):
    headers = [clean(ws.cell(1, col).value) for col in range(1, min(ws.max_column, 18) + 1)]
    print(f"  - {ws.title}: rows={ws.max_row}, cols={ws.max_column}, images={len(getattr(ws, '_images', []))}")
    print(f"    headers: {headers}")
    for row_idx in range(2, min(ws.max_row, 6) + 1):
        row = [clean(ws.cell(row_idx, col).value) for col in range(1, min(ws.max_column, 14) + 1)]
        if any(row):
            print(f"    row {row_idx}: {row}")
    formulas = formula_samples(ws)
    print(f"    formulas: {len(formulas)} samples")
    for item in formulas:
        print(f"      {item}")
    links, samples = hyperlink_count(ws)
    print(f"    hyperlinks: {links}")
    for item in samples:
        print(f"      {item}")


def main():
    for name in FILES:
        print(f"\n{name}")
        wb = openpyxl.load_workbook(BASE / name, data_only=False)
        for ws in wb.worksheets:
            print_sheet(ws)


if __name__ == "__main__":
    main()
