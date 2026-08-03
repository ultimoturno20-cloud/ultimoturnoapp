from pathlib import Path
import json
import openpyxl


BASE = Path(r"C:\Users\skype\Downloads\New folder")
FILES = [
    "Presupuestador.xlsx",
    "extraer imagen.xlsx",
    "Generar imagen adelanto.xlsx",
    "Claim 12 de Julio.xlsx",
]


def cell_text(cell):
    value = cell.value
    if value is None:
        value = ""
    text = str(value)
    if len(text) > 120:
        text = text[:117] + "..."
    return text


def summarize_sheet(ws):
    rows = ws.max_row
    cols = ws.max_column
    non_empty_rows = 0
    formulas = []
    hyperlinks = []

    for row in ws.iter_rows():
        has_value = False
        for cell in row:
            if cell.value not in (None, ""):
                has_value = True
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append({"cell": cell.coordinate, "formula": cell_text(cell)})
                if cell.hyperlink:
                    hyperlinks.append({"cell": cell.coordinate, "target": str(cell.hyperlink.target)[:160]})
        if has_value:
            non_empty_rows += 1

    preview_rows = min(rows, 12)
    preview_cols = min(cols, 16)
    preview = []
    for r in range(1, preview_rows + 1):
        preview.append([cell_text(ws.cell(r, c)) for c in range(1, preview_cols + 1)])

    return {
        "title": ws.title,
        "max_row": rows,
        "max_col": cols,
        "non_empty_rows": non_empty_rows,
        "merged_ranges": [str(m) for m in list(ws.merged_cells.ranges)[:20]],
        "formula_count": len(formulas),
        "formula_samples": formulas[:20],
        "hyperlink_count": len(hyperlinks),
        "hyperlink_samples": hyperlinks[:20],
        "image_count": len(getattr(ws, "_images", [])),
        "preview": preview,
    }


def main():
    report = {}
    for name in FILES:
        path = BASE / name
        wb = openpyxl.load_workbook(path, data_only=False)
        report[name] = {
            "sheets": [summarize_sheet(ws) for ws in wb.worksheets],
        }
    print(json.dumps(report, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
